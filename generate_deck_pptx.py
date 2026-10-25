#!/usr/bin/env python3
"""
Generate the Reach Pitch Deck in PowerPoint (.pptx) format.
Clinical Cloudflare-inspired product announcement aesthetic matching Bridge.pptx.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ── Color Palette ─────────────────────────────────────────────────────────────
DARK_BG = RGBColor(0x0F, 0x17, 0x2A)       # Slate 900
LIGHT_BG = RGBColor(0xFF, 0xFF, 0xFF)      # White
ORANGE = RGBColor(0xEA, 0x58, 0x0C)        # Accent Orange
ORANGE_LIGHT = RGBColor(0xF9, 0x73, 0x16)  # Highlights
DARK_TEXT = RGBColor(0x0F, 0x17, 0x2A)     # Near-black
WHITE_TEXT = RGBColor(0xFF, 0xFF, 0xFF)    # Pure White
SLATE_500 = RGBColor(0x64, 0x74, 0x8B)     # Muted body (light)
SLATE_400 = RGBColor(0x94, 0xA3, 0xB8)     # Muted body (dark)
SLATE_300 = RGBColor(0xCB, 0xD5, 0xE1)     # Subdued
RED_TEXT = RGBColor(0x99, 0x1B, 0x1B)      # Red accent
GREEN_TEXT = RGBColor(0x16, 0x65, 0x34)    # Green accent

CARD_LIGHT_BG = RGBColor(0xF8, 0xFA, 0xFC) # Slate 50
CARD_LIGHT_BORDER = RGBColor(0xE2, 0xE8, 0xF0) # Slate 200
CARD_DARK_BG = RGBColor(0x1E, 0x29, 0x3B)  # Slate 800
CARD_DARK_BORDER = RGBColor(0x33, 0x41, 0x55)  # Slate 700

# ── Dimensions (16:9 Widescreen) ─────────────────────────────────────────────
SLIDE_WIDTH = 12192000   # 13.333 inches
SLIDE_HEIGHT = 6858000   # 7.5 inches

TOTAL_SLIDES = 10

def create_deck():
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    blank_layout = prs.slide_layouts[6]

    def base_slide(is_dark=False, slide_idx=1, notes=""):
        slide = prs.slides.add_slide(blank_layout)
        
        # Background
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT)
        bg.fill.solid()
        bg.fill.fore_color.rgb = DARK_BG if is_dark else LIGHT_BG
        bg.line.fill.background()

        # Bottom thin progress line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, SLIDE_HEIGHT - 38100, int(SLIDE_WIDTH * (slide_idx / TOTAL_SLIDES)), 38100)
        line.fill.solid()
        line.fill.fore_color.rgb = ORANGE
        line.line.fill.background()

        # Slide counter at bottom right
        counter_box = slide.shapes.add_textbox(SLIDE_WIDTH - Inches(1.8), SLIDE_HEIGHT - Inches(0.45), Inches(1.5), Inches(0.3))
        tf = counter_box.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = f"{slide_idx:02d} / {TOTAL_SLIDES:02d}"
        r.font.name = "JetBrains Mono"
        r.font.size = Pt(9.5)
        r.font.color.rgb = SLATE_400 if is_dark else SLATE_500

        # Speaker notes
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

        return slide

    def add_header(slide, kicker, headline, is_dark=False):
        # Kicker
        kicker_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.5), Inches(0.3))
        tf = kicker_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = kicker.upper()
        r.font.name = "JetBrains Mono"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = ORANGE

        # Headline
        head_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.95), Inches(11.5), Inches(0.6))
        tf = head_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = headline
        r.font.name = "Inter"
        r.font.size = Pt(26)
        r.font.bold = True
        r.font.color.rgb = WHITE_TEXT if is_dark else DARK_TEXT

    def add_footer(slide, summary, is_dark=False):
        foot_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.55), Inches(10.5), Inches(0.35))
        tf = foot_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = summary
        r.font.name = "Inter"
        r.font.size = Pt(11)
        r.font.color.rgb = SLATE_400 if is_dark else SLATE_500

    def add_card(slide, left, top, width, height, is_dark=False, border_color=None, fill_color=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = fill_color if fill_color else (CARD_DARK_BG if is_dark else CARD_LIGHT_BG)
        card.line.color.rgb = border_color if border_color else (CARD_DARK_BORDER if is_dark else CARD_LIGHT_BORDER)
        card.line.width = Pt(1)
        return card

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 1: COVER (Dark)
    # ═════════════════════════════════════════════════════════════════════════
    s1 = base_slide(is_dark=True, slide_idx=1, notes="Open cold. Name the company and the product. Pause on REACH.")
    
    # Kicker
    k1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(10), Inches(0.4))
    tf = k1.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "LOCAL SERVICES & ESCROW"
    r.font.name = "JetBrains Mono"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = ORANGE

    # Display Title
    t1 = s1.shapes.add_textbox(Inches(0.95), Inches(2.2), Inches(10), Inches(1.5))
    tf = t1.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "REACH"
    r.font.name = "Inter"
    r.font.size = Pt(72)
    r.font.bold = True
    r.font.color.rgb = WHITE_TEXT

    # Subtitle
    sub1 = s1.shapes.add_textbox(Inches(1.0), Inches(3.7), Inches(10), Inches(0.6))
    tf = sub1.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Find anything, close to home."
    r.font.name = "Inter"
    r.font.size = Pt(22)
    r.font.color.rgb = SLATE_400

    # Pill tags
    tags = ["SOLANA ESCROW", "FIAT RAILS", "POSTGIS MATCHING"]
    tag_left = Inches(1.0)
    for tag in tags:
        pill = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, tag_left, Inches(4.5), Inches(2.2), Inches(0.4))
        pill.fill.solid()
        pill.fill.fore_color.rgb = CARD_DARK_BG
        pill.line.color.rgb = CARD_DARK_BORDER
        tf = pill.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = tag
        r.font.name = "JetBrains Mono"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = ORANGE if tag == "SOLANA ESCROW" else SLATE_300
        tag_left += Inches(2.4)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 2: THE TRUST GAP (Light)
    # ═════════════════════════════════════════════════════════════════════════
    s2 = base_slide(is_dark=False, slide_idx=2, notes="Frame the pain: providers and clients both carry risk. The status quo has no guarantees.")
    add_header(s2, "THE TRUST GAP", "Money Moves on Hope. Not Guarantees.")
    add_footer(s2, "A ₦4T+ informal services market running entirely on social trust and informal referrals.")

    # Left card: Status Quo
    add_card(s2, Inches(0.8), Inches(1.75), Inches(5.6), Inches(4.55))
    tb_left = s2.shapes.add_textbox(Inches(1.0), Inches(1.95), Inches(5.2), Inches(4.15))
    tf_l = tb_left.text_frame
    tf_l.word_wrap = True
    
    p = tf_l.paragraphs[0]
    r = p.add_run()
    r.text = "STATUS QUO\n"
    r.font.name = "JetBrains Mono"
    r.font.size = Pt(9.5)
    r.font.bold = True
    r.font.color.rgb = RED_TEXT

    p2 = tf_l.add_paragraph()
    r = p2.add_run()
    r.text = "Local Services Today\n"
    r.font.name = "Inter"
    r.font.size = Pt(17)
    r.font.bold = True
    r.font.color.rgb = RED_TEXT

    pain_points = [
        ("01", "No Discovery Hub", "No single platform to find nearby specialists, dispatch errands, or rent an extra set of hands."),
        ("02", "No Payment Guarantee", "Providers complete physical work with zero contractual assurance of getting paid."),
        ("03", "Upfront Client Risk", "Clients pay before seeing results, with zero recourse or protection if work fails."),
        ("04", "Word-of-Mouth Pricing", "Arbitrary prices with no reliable benchmark or transparent market rates.")
    ]

    for num, title, desc in pain_points:
        p = tf_l.add_paragraph()
        p.space_before = Pt(8)
        r1 = p.add_run()
        r1.text = f"{num}  "
        r1.font.name = "JetBrains Mono"
        r1.font.size = Pt(10)
        r1.font.bold = True
        r1.font.color.rgb = ORANGE

        r2 = p.add_run()
        r2.text = f"{title} — "
        r2.font.name = "Inter"
        r2.font.size = Pt(10.5)
        r2.font.bold = True
        r2.font.color.rgb = DARK_TEXT

        r3 = p.add_run()
        r3.text = desc
        r3.font.name = "Inter"
        r3.font.size = Pt(10)
        r3.font.color.rgb = SLATE_500

    # Right card: The Structural Gap
    add_card(s2, Inches(6.7), Inches(1.75), Inches(5.6), Inches(4.55))
    tb_right = s2.shapes.add_textbox(Inches(6.9), Inches(1.95), Inches(5.2), Inches(4.15))
    tf_r = tb_right.text_frame
    tf_r.word_wrap = True

    p = tf_r.paragraphs[0]
    r = p.add_run()
    r.text = "THE STRUCTURAL GAP\n"
    r.font.name = "JetBrains Mono"
    r.font.size = Pt(9.5)
    r.font.bold = True
    r.font.color.rgb = SLATE_500

    p2 = tf_r.add_paragraph()
    r = p2.add_run()
    r.text = "Why Existing Solutions Fail\n"
    r.font.name = "Inter"
    r.font.size = Pt(17)
    r.font.bold = True
    r.font.color.rgb = DARK_TEXT

    gap_bullets = [
        "No verified nearby specialist directory or location radius.",
        "No escrow contracts, no collateral deposits, no legal recourse.",
        "Trust is fragmented across WhatsApp groups and unrecorded calls.",
        "A ₦4T+ informal economy operating entirely on hope and social friction."
    ]

    for pt in gap_bullets:
        p = tf_r.add_paragraph()
        p.space_before = Pt(14)
        r1 = p.add_run()
        r1.text = "→  "
        r1.font.name = "JetBrains Mono"
        r1.font.size = Pt(11)
        r1.font.bold = True
        r1.font.color.rgb = ORANGE

        r2 = p.add_run()
        r2.text = pt
        r2.font.name = "Inter"
        r2.font.size = Pt(11)
        r2.font.color.rgb = SLATE_500

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 3: THE SOLUTION (Light)
    # ═════════════════════════════════════════════════════════════════════════
    s3 = base_slide(is_dark=False, slide_idx=3, notes="Three steps. The user never sees the chain — they just feel the guarantee.")
    add_header(s3, "ESCROW-POWERED SERVICES", "Make Every Transaction Trustless.")
    add_footer(s3, "Blockchain settles the transaction cryptographically. Ordinary users never see a wallet or private key.")

    steps = [
        ("01", "TASK CREATION", "Funds Locked in Escrow", "Client's budget is locked in a Solana escrow PDA before provider mobilizes. Provider has verifiable proof funds exist."),
        ("02", "SETTLEMENT GATE", "Work Verified & Approved", "Provider completes task and uploads proof. Client reviews, approves, and marks the contract fulfilled."),
        ("03", "TRUSTLESS PAYOUT", "Automated Payout Release", "Smart contract automatically releases USDC payout to provider. In disputes, admin mediates with custom on-chain splits.")
    ]

    c_left = Inches(0.8)
    for num, badge, title, desc in steps:
        add_card(s3, c_left, Inches(1.85), Inches(3.6), Inches(4.35))
        tb = s3.shapes.add_textbox(c_left + Inches(0.2), Inches(2.05), Inches(3.2), Inches(3.95))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        r1 = p.add_run()
        r1.text = f"{num} · "
        r1.font.name = "JetBrains Mono"
        r1.font.size = Pt(10)
        r1.font.bold = True
        r1.font.color.rgb = ORANGE

        r2 = p.add_run()
        r2.text = badge
        r2.font.name = "JetBrains Mono"
        r2.font.size = Pt(9.5)
        r2.font.bold = True
        r2.font.color.rgb = SLATE_500

        p_t = tf.add_paragraph()
        p_t.space_before = Pt(14)
        r = p_t.add_run()
        r.text = title
        r.font.name = "Inter"
        r.font.size = Pt(17)
        r.font.bold = True
        r.font.color.rgb = DARK_TEXT

        p_d = tf.add_paragraph()
        p_d.space_before = Pt(10)
        r = p_d.add_run()
        r.text = desc
        r.font.name = "Inter"
        r.font.size = Pt(11)
        r.font.color.rgb = SLATE_500

        c_left += Inches(3.9)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 4: HOW REACH WORKS (Light)
    # ═════════════════════════════════════════════════════════════════════════
    s4 = base_slide(is_dark=False, slide_idx=4, notes="Three zones. The middle zone is where Reach earns its keep — escrow plus geospatial matching.")
    add_header(s4, "HOW REACH WORKS", "Simple for Users. Powerful Underneath.")
    add_footer(s4, "Three funding rails: Naira cards, bank transfers, or native Solana wallet (Phantom / Solflare).")

    zones = [
        ("01", "NAIRA RAILS", "Post & Fund in Naira", "Client posts a task with budget and location radius. Pays instantly via Nigerian debit card or instant bank transfer.", "No seed phrases. No crypto onboarding. Fiat-first simplicity."),
        ("02", "CUSTODIAL RELAYER", "Escrow & Matchmaking", "Server-side relayer signs and deposits USDC into Solana PDA. PostGIS engine alerts verified nearby specialists.", "Dual-custody float management with Obiex & FossaPay integration."),
        ("03", "USDC / NGN OFFRAMP", "Work, Deliver & Cash Out", "Provider accepts task, requests optional advance for materials, completes work, and receives instant payment.", "Receives USDC directly or off-ramps straight to Nigerian bank accounts.")
    ]

    c_left = Inches(0.8)
    for num, badge, title, desc, sub in zones:
        add_card(s4, c_left, Inches(1.85), Inches(3.6), Inches(4.35))
        tb = s4.shapes.add_textbox(c_left + Inches(0.2), Inches(2.05), Inches(3.2), Inches(3.95))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        r1 = p.add_run()
        r1.text = f"{num} · "
        r1.font.name = "JetBrains Mono"
        r1.font.size = Pt(10)
        r1.font.bold = True
        r1.font.color.rgb = ORANGE

        r2 = p.add_run()
        r2.text = badge
        r2.font.name = "JetBrains Mono"
        r2.font.size = Pt(9.5)
        r2.font.bold = True
        r2.font.color.rgb = ORANGE if "RELAYER" in badge else SLATE_500

        p_t = tf.add_paragraph()
        p_t.space_before = Pt(14)
        r = p_t.add_run()
        r.text = title
        r.font.name = "Inter"
        r.font.size = Pt(17)
        r.font.bold = True
        r.font.color.rgb = DARK_TEXT

        p_d = tf.add_paragraph()
        p_d.space_before = Pt(10)
        r = p_d.add_run()
        r.text = desc
        r.font.name = "Inter"
        r.font.size = Pt(11)
        r.font.color.rgb = SLATE_500

        p_s = tf.add_paragraph()
        p_s.space_before = Pt(16)
        r = p_s.add_run()
        r.text = sub
        r.font.name = "Inter"
        r.font.size = Pt(10)
        r.font.italic = True
        r.font.color.rgb = SLATE_400

        c_left += Inches(3.9)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 5: THE PLATFORM (Light)
    # ═════════════════════════════════════════════════════════════════════════
    s5 = base_slide(is_dark=False, slide_idx=5, notes="Six features, shipped. Emphasize this is production code, not a prototype.")
    add_header(s5, "THE PLATFORM", "What We've Built.")
    add_footer(s5, "PWA-ready. Push notifications via Firebase. Cloudinary media pipeline. Production engineering.")

    features = [
        ("01", "Geospatial Matching", "PostGIS-powered engine matches tasks with verified providers within dynamic km radiuses."),
        ("02", "On-Chain Escrow", "Solana Anchor program handling collateral deposits, upfront advances, and automated release."),
        ("03", "Fiat Payment Rails", "Native integration with Obiex and FossaPay for local card checkout and bank off-ramps."),
        ("04", "Real-Time Messaging", "In-app client-provider chat for instant coordination, media sharing, and milestone updates."),
        ("05", "Reputation System", "Bidirectional reviews and verified completion metrics building permanent on-chain trust."),
        ("06", "Identity & KYC", "National Identity Number (NIN) and Bank Verification Number (BVN) checks via automated APIs.")
    ]

    for idx, (num, title, desc) in enumerate(features):
        row = idx // 3
        col = idx % 3
        c_left = Inches(0.8) + col * Inches(3.9)
        c_top = Inches(1.85) + row * Inches(2.25)
        add_card(s5, c_left, c_top, Inches(3.6), Inches(2.05))

        tb = s5.shapes.add_textbox(c_left + Inches(0.2), c_top + Inches(0.18), Inches(3.2), Inches(1.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        r1 = p.add_run()
        r1.text = f"{num}  "
        r1.font.name = "JetBrains Mono"
        r1.font.size = Pt(10)
        r1.font.bold = True
        r1.font.color.rgb = ORANGE

        r2 = p.add_run()
        r2.text = title
        r2.font.name = "Inter"
        r2.font.size = Pt(14)
        r2.font.bold = True
        r2.font.color.rgb = DARK_TEXT

        p_d = tf.add_paragraph()
        p_d.space_before = Pt(6)
        r = p_d.add_run()
        r.text = desc
        r.font.name = "Inter"
        r.font.size = Pt(10)
        r.font.color.rgb = SLATE_500

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 6: MARKET OPPORTUNITY (Light)
    # ═════════════════════════════════════════════════════════════════════════
    s6 = base_slide(is_dark=False, slide_idx=6, notes="The market is young, mobile-first, and crypto-fluent. Akwa Ibom (Uyo) first, then the region.")
    add_header(s6, "MARKET OPPORTUNITY", "Nigeria's Service Economy Is Massive and Underserved.")
    add_footer(s6, "Mobile-first, young, and crypto-fluent. High inflation drives everyday demand for stablecoin earnings.")

    stats = [
        ("200M+", "Population", "Fastest-growing urban market in Africa"),
        ("70%", "Under 30", "Mobile-first, digital native workforce"),
        ("#1", "Crypto Adoption", "Chainalysis Global Adoption Index (Africa)"),
        ("₦4T+", "Informal Services", "Unorganized, cash-heavy addressable TAM")
    ]

    s_left = Inches(0.8)
    for val, lbl, hint in stats:
        add_card(s6, s_left, Inches(1.85), Inches(2.65), Inches(2.0))
        tb = s6.shapes.add_textbox(s_left + Inches(0.15), Inches(1.95), Inches(2.35), Inches(1.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p_v = tf.paragraphs[0]
        r = p_v.add_run()
        r.text = val
        r.font.name = "Inter"
        r.font.size = Pt(32)
        r.font.bold = True
        r.font.color.rgb = ORANGE

        p_l = tf.add_paragraph()
        p_l.space_before = Pt(2)
        r = p_l.add_run()
        r.text = lbl
        r.font.name = "Inter"
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = DARK_TEXT

        p_h = tf.add_paragraph()
        p_h.space_before = Pt(2)
        r = p_h.add_run()
        r.text = hint
        r.font.name = "Inter"
        r.font.size = Pt(9.5)
        r.font.color.rgb = SLATE_500

        s_left += Inches(2.95)

    # 2 Rollout cards below
    add_card(s6, Inches(0.8), Inches(4.1), Inches(5.6), Inches(2.2))
    tb_u = s6.shapes.add_textbox(Inches(1.0), Inches(4.25), Inches(5.2), Inches(1.9))
    tf_u = tb_u.text_frame
    tf_u.word_wrap = True

    p = tf_u.paragraphs[0]
    r = p.add_run()
    r.text = "PHASE 1 LAUNCH\n"
    r.font.name = "JetBrains Mono"
    r.font.size = Pt(9.5)
    r.font.bold = True
    r.font.color.rgb = ORANGE

    p2 = tf_u.add_paragraph()
    r = p2.add_run()
    r.text = "Akwa Ibom & Uyo Launchpad\n"
    r.font.name = "Inter"
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = DARK_TEXT

    p3 = tf_u.add_paragraph()
    r = p3.add_run()
    r.text = "Emerging digital hub with high artisan density, strong smartphone penetration, and an active, tight-knit service economy."
    r.font.name = "Inter"
    r.font.size = Pt(11)
    r.font.color.rgb = SLATE_500

    add_card(s6, Inches(6.7), Inches(4.1), Inches(5.6), Inches(2.2))
    tb_e = s6.shapes.add_textbox(Inches(6.9), Inches(4.25), Inches(5.2), Inches(1.9))
    tf_e = tb_e.text_frame
    tf_e.word_wrap = True

    p = tf_e.paragraphs[0]
    r = p.add_run()
    r.text = "EXPANSION PATH\n"
    r.font.name = "JetBrains Mono"
    r.font.size = Pt(9.5)
    r.font.bold = True
    r.font.color.rgb = SLATE_500

    p2 = tf_e.add_paragraph()
    r = p2.add_run()
    r.text = "Regional & National Rollout\n"
    r.font.name = "Inter"
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = DARK_TEXT

    p3 = tf_e.add_paragraph()
    r = p3.add_run()
    r.text = "Scale along the South-South corridor (Port Harcourt, Calabar) to Lagos and Abuja, followed by regional West Africa."
    r.font.name = "Inter"
    r.font.size = Pt(11)
    r.font.color.rgb = SLATE_500

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 7: REVENUE MODEL (Light)
    # ═════════════════════════════════════════════════════════════════════════
    s7 = base_slide(is_dark=False, slide_idx=7, notes="Three streams, all on-chain and auditable. Revenue is protocol infrastructure, not an opaque ledger.")
    add_header(s7, "REVENUE MODEL", "Transparent, On-Chain Revenue.")
    add_footer(s7, "Protocol revenue is enforced by smart contract logic and accumulates in a publicly auditable fee vault.")

    rev_streams = [
        ("2.5%", "BPS: 250", "Task Creation Fee", "Charged upfront when escrow PDA is funded. Diverts directly into the on-chain protocol fee vault."),
        ("5.0%", "BPS: 500", "Platform Commission", "Deducted from the approved provider payout upon successful completion of the assignment."),
        ("10.0%", "BPS: 1000", "Cancellation Fee", "Charged if creator cancels an active task after provider assignment to protect worker time.")
    ]

    r_left = Inches(0.8)
    for pct, bps, title, desc in rev_streams:
        add_card(s7, r_left, Inches(1.85), Inches(3.6), Inches(2.9))
        tb = s7.shapes.add_textbox(r_left + Inches(0.2), Inches(2.0), Inches(3.2), Inches(2.6))
        tf = tb.text_frame
        tf.word_wrap = True

        p_b = tf.paragraphs[0]
        r = p_b.add_run()
        r.text = bps
        r.font.name = "JetBrains Mono"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = ORANGE

        p_p = tf.add_paragraph()
        p_p.space_before = Pt(4)
        r = p_p.add_run()
        r.text = pct
        r.font.name = "Inter"
        r.font.size = Pt(36)
        r.font.bold = True
        r.font.color.rgb = ORANGE

        p_t = tf.add_paragraph()
        r = p_t.add_run()
        r.text = title
        r.font.name = "Inter"
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = DARK_TEXT

        p_d = tf.add_paragraph()
        p_d.space_before = Pt(4)
        r = p_d.add_run()
        r.text = desc
        r.font.name = "Inter"
        r.font.size = Pt(10)
        r.font.color.rgb = SLATE_500

        r_left += Inches(3.9)

    # Full-width Callout Card
    add_card(s7, Inches(0.8), Inches(4.95), Inches(11.5), Inches(1.3), is_dark=True, fill_color=DARK_BG, border_color=CARD_DARK_BORDER)
    tb_c = s7.shapes.add_textbox(Inches(1.1), Inches(5.1), Inches(10.9), Inches(1.0))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True

    p = tf_c.paragraphs[0]
    r = p.add_run()
    r.text = "VERIFIABLE BALANCE SHEET  ·  SOLANA MAINNET PDA\n"
    r.font.name = "JetBrains Mono"
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = ORANGE

    p2 = tf_c.add_paragraph()
    r = p2.add_run()
    r.text = "All protocol fees collect directly into the Anchor Fee Vault PDA. Platform revenue is public, mathematically guaranteed, and immune to opaque bookkeeping."
    r.font.name = "Inter"
    r.font.size = Pt(12)
    r.font.color.rgb = WHITE_TEXT

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 8: TECHNICAL ARCHITECTURE (Dark)
    # ═════════════════════════════════════════════════════════════════════════
    s8 = base_slide(is_dark=True, slide_idx=8, notes="Five layers, all production. Emphasize the server-side Solana relayer — users never touch keys.")
    add_header(s8, "TECHNICAL ARCHITECTURE", "Built for Scale and Trust.", is_dark=True)
    add_footer(s8, "158 backend tests passing. Dual-custody treasury management. Production-ready codebase.", is_dark=True)

    arch_layers = [
        ("01", "FRONTEND", "Next.js 16 + React 19 + TypeScript + Tailwind v4", "PWA-enabled, responsive mobile interface with end-to-end type safety."),
        ("02", "BACKEND", "FastAPI (async) + SQLAlchemy 2.0 + asyncpg + Pydantic v2", "High-throughput async Python service with custodial transaction engine."),
        ("03", "DATA PLANE", "PostgreSQL + PostGIS + Redis Cache", "Spatial indexing for millisecond-level nearby matching and query acceleration."),
        ("04", "ESCROW CONTRACT", "Solana Anchor Program (Rust, 2,300+ LOC)", "Canonical escrow, provider collateral, debt tracking, and dispute splits."),
        ("05", "INTEGRATIONS", "Obiex API + FossaPay + Firebase Push + Cloudinary", "Naira checkout, bank payouts, biometric KYC, and instant dispatch notifications.")
    ]

    a_top = Inches(1.8)
    for num, layer, tech, detail in arch_layers:
        card = add_card(s8, Inches(0.8), a_top, Inches(11.5), Inches(0.85), is_dark=True)
        tb = s8.shapes.add_textbox(Inches(1.0), a_top + Inches(0.08), Inches(11.1), Inches(0.68))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        r1 = p.add_run()
        r1.text = f"{num} · {layer}   "
        r1.font.name = "JetBrains Mono"
        r1.font.size = Pt(10)
        r1.font.bold = True
        r1.font.color.rgb = ORANGE

        r2 = p.add_run()
        r2.text = f"{tech}   "
        r2.font.name = "Inter"
        r2.font.size = Pt(12)
        r2.font.bold = True
        r2.font.color.rgb = WHITE_TEXT

        r3 = p.add_run()
        r3.text = f"—  {detail}"
        r3.font.name = "Inter"
        r3.font.size = Pt(10.5)
        r3.font.color.rgb = SLATE_400

        a_top += Inches(0.92)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 9: EXECUTION TIMELINE (Light)
    # ═════════════════════════════════════════════════════════════════════════
    s9 = base_slide(is_dark=False, slide_idx=9, notes="Backend is built, contract is audited. We're closing the last mile to market.")
    add_header(s9, "EXECUTION TIMELINE", "From Code to Market.")
    add_footer(s9, "The backend is built. The smart contract is audited. We are closing the last mile to commercial launch.")

    milestones = [
        ("Q3 2026", "COMPLETE", GREEN_TEXT, "Backend & Contracts", "Solana Anchor program audited & patched. FastAPI relayer & treasury engine built with 158 tests passing."),
        ("Q4 2026", "IN FLIGHT", ORANGE, "Staging & Integration", "Finalizing checkout screens, wallet connection modals, and live Obiex business settlement tests."),
        ("Q1 2027", "TARGET", SLATE_500, "Uyo Beta Launch", "Controlled beta release across Uyo commercial clusters. First 1,000 live escrow task settlements."),
        ("Q2 2027", "TARGET", SLATE_500, "Multi-City Scale", "Physical trade items added. Expansion to Port Harcourt, Calabar, and Lagos."),
        ("Q3 2027", "PLANNED", SLATE_500, "Multichain & Regional", "NEAR Intents cross-chain integration. Entering regional West African commercial centers (Accra).")
    ]

    m_left = Inches(0.8)
    for quarter, badge, bcolor, title, desc in milestones:
        add_card(s9, m_left, Inches(1.85), Inches(2.2), Inches(4.35))
        tb = s9.shapes.add_textbox(m_left + Inches(0.12), Inches(2.05), Inches(1.96), Inches(3.95))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = quarter
        r.font.name = "JetBrains Mono"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = ORANGE

        p_b = tf.add_paragraph()
        p_b.space_before = Pt(4)
        r = p_b.add_run()
        r.text = badge
        r.font.name = "JetBrains Mono"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = bcolor

        p_t = tf.add_paragraph()
        p_t.space_before = Pt(12)
        r = p_t.add_run()
        r.text = title
        r.font.name = "Inter"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = DARK_TEXT

        p_d = tf.add_paragraph()
        p_d.space_before = Pt(8)
        r = p_d.add_run()
        r.text = desc
        r.font.name = "Inter"
        r.font.size = Pt(9.5)
        r.font.color.rgb = SLATE_500

        m_left += Inches(2.33)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 10: CLOSING (Dark)
    # ═════════════════════════════════════════════════════════════════════════
    s10 = base_slide(is_dark=True, slide_idx=10, notes="Close on the three-word promise. Leave the contact on screen.")
    
    # Center Kicker
    k10 = s10.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(10.2), Inches(0.4))
    tf = k10.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "REACH"
    r.font.name = "JetBrains Mono"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = ORANGE

    # Big 3-Line Statement
    t10 = s10.shapes.add_textbox(Inches(1.0), Inches(2.3), Inches(10.2), Inches(2.2))
    tf = t10.text_frame
    
    p1 = tf.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    r1 = p1.add_run()
    r1.text = "YOUR MONEY."
    r1.font.name = "Inter"
    r1.font.size = Pt(54)
    r1.font.bold = True
    r1.font.color.rgb = WHITE_TEXT

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = "YOUR WORK."
    r2.font.name = "Inter"
    r2.font.size = Pt(54)
    r2.font.bold = True
    r2.font.color.rgb = WHITE_TEXT

    p3 = tf.add_paragraph()
    p3.alignment = PP_ALIGN.CENTER
    r3 = p3.add_run()
    r3.text = "SETTLED."
    r3.font.name = "Inter"
    r3.font.size = Pt(54)
    r3.font.bold = True
    r3.font.color.rgb = ORANGE_LIGHT

    # Subhead
    sub10 = s10.shapes.add_textbox(Inches(1.0), Inches(4.7), Inches(10.2), Inches(0.5))
    tf = sub10.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Trustless escrow for Africa's local service economy."
    r.font.name = "Inter"
    r.font.size = Pt(17)
    r.font.color.rgb = SLATE_400

    # Thin Divider
    div = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.8), Inches(5.4), Inches(0.6), Inches(0.02))
    div.fill.solid()
    div.fill.fore_color.rgb = CARD_DARK_BORDER
    div.line.fill.background()

    # Contact Info
    c10 = s10.shapes.add_textbox(Inches(1.0), Inches(5.65), Inches(10.2), Inches(0.5))
    tf = c10.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    
    r_dom = p.add_run()
    r_dom.text = "reachapp.xyz"
    r_dom.font.name = "JetBrains Mono"
    r_dom.font.size = Pt(14)
    r_dom.font.bold = True
    r_dom.font.color.rgb = WHITE_TEXT

    r_sep = p.add_run()
    r_sep.text = "   ·   "
    r_sep.font.name = "JetBrains Mono"
    r_sep.font.size = Pt(14)
    r_sep.font.color.rgb = SLATE_500

    r_mail = p.add_run()
    r_mail.text = "hello@peni.dev"
    r_mail.font.name = "JetBrains Mono"
    r_mail.font.size = Pt(14)
    r_mail.font.bold = True
    r_mail.font.color.rgb = ORANGE

    output_path = "/home/peni/Projects/reach/Reach - Pitch Deck.pptx"
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    create_deck()
